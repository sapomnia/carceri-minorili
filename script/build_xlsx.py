"""Costruisce carceri_minorili_ipm_2025_2026.xlsx dai CSV estratti (regole riccardo-xlsx)."""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

OUT_DIR, OUT_2025, GEO_DIR, DEST = sys.argv[1:5]
sys.path.insert(0, GEO_DIR)
from geo_lookup import lookup_codice  # noqa: E402

COMUNE = {'Torino': 'Torino', 'Pontremoli (MS)': 'Pontremoli', 'Milano': 'Milano', 'Rovigo': 'Rovigo',
          'Treviso': 'Treviso', 'Bologna': 'Bologna', 'Bologna Sezione Dozza': 'Bologna', 'Firenze': 'Firenze',
          'Roma': 'Roma', "L'Aquila": "L'Aquila", 'Lecce': 'Lecce', 'Nisida (NA)': 'Napoli',
          'Airola (BN)': 'Airola', 'Bari': 'Bari', 'Potenza': 'Potenza', 'Catanzaro': 'Catanzaro',
          'Palermo': 'Palermo', 'Catania': 'Catania', 'Acireale (CT)': 'Acireale', 'Caltanissetta': 'Caltanissetta',
          'Quartucciu (CA)': 'Quartucciu'}
NOME = {'Torino': 'Torino "Ferrante Aporti"', 'Milano': 'Milano "Cesare Beccaria"', 'Roma': 'Roma "Casal del Marmo"',
        'Nisida (NA)': 'Nisida (NA)', 'Bologna': 'Bologna "Pratello"'}
ORDINE = list(COMUNE)

# Capienza: DGMC non la pubblica nei report statistici. Fonti: risposta del ministro all'interrogazione 4-06110
# (18/11/2025), Carta dei servizi 2026 IPM Cagliari, schede istituto dell'Osservatorio Antigone (ragazzidentro.it).
ANT = 'Antigone, scheda istituto (ragazzidentro.it)'
CAP = {
    'Torino': (46, ANT, ''),
    'Pontremoli (MS)': (17, ANT, 'Solo femminile'),
    'Milano': (46, ANT, 'Antigone indica 46/48 posti'),
    'Rovigo': (22, 'Ministero della Giustizia, risposta interrogazione 4-06110 (18/11/2025)',
               'Più 8 posti di custodia attenuata; attivo dal 23/02/2026'),
    'Treviso': (12, ANT, 'Attività sospesa dal 23/02/2026'),
    'Bologna': (44, ANT, ''),
    'Firenze': (17, 'Ministero della Giustizia, risposta interrogazione 4-06110 (18/11/2025)',
                'Lavori in corso per salire a 29 posti'),
    'Roma': (52, ANT, 'Solo palazzine maschili (25 giovani adulti, 27 minorenni); posti della sezione femminile non indicati'),
    "L'Aquila": (28, 'Ministero della Giustizia, risposta interrogazione 4-06110 (18/11/2025)', 'Attivo dal 28/10/2025'),
    'Lecce': (27, 'Ministero della Giustizia, risposta interrogazione 4-06110 (18/11/2025)', 'Attivo dal 28/12/2025'),
    'Nisida (NA)': (74, ANT, 'Capienza ridotta da lavori e spazi inagibili alla visita'),
    'Airola (BN)': (42, ANT, 'Ridotta a 18 posti durante i lavori'),
    'Bari': (35, ANT, ''),
    'Potenza': (17, ANT, ''),
    'Catanzaro': (35, ANT, '33 posti disponibili alla visita (cella inagibile)'),
    'Palermo': (40, ANT, '28 posti disponibili alla visita (lavori)'),
    'Catania': (52, ANT, '43 posti effettivi alla visita (9 celle inagibili)'),
    'Acireale (CT)': (20, ANT, 'Capienza effettiva indicata da Antigone'),
    'Caltanissetta': (12, ANT, '8 posti alla visita (stanza inagibile)'),
    'Quartucciu (CA)': (10, 'Carta dei servizi 2026 IPM e CPA di Cagliari', 'Ridotta di fatto a 8; ampliamento in corso'),
}

# Capienza DGMC al 30/03/2025 dal report del Garante nazionale (tabelle 1 e 2, maschi + femmine): tiene conto di lavori
# e camere inagibili. È la capienza usata per l'affollamento; gli IPM aperti dopo restano sulle fonti di CAP.
GARANTE = 'Garante nazionale, Osservatorio penitenziario minori, report al 31/03/2025 (dati DGMC al 30/03/2025)'
CAP_DGMC = {'Acireale (CT)': 20, 'Airola (BN)': 24, 'Bari': 30, 'Bologna': 40, 'Caltanissetta': 12, 'Catania': 46,
            'Catanzaro': 32, 'Firenze': 17, 'Milano': 48, 'Nisida (NA)': 76, 'Palermo': 28, 'Potenza': 16,
            'Quartucciu (CA)': 10, 'Roma': 65, 'Torino': 46, 'Treviso': 12, 'Pontremoli (MS)': 16}
NOTE_DGMC = {'Roma': '52 posti maschili e 13 femminili', 'Pontremoli (MS)': 'Solo femminile',
             'Treviso': 'Attività sospesa dal 23/02/2026'}
# Presenti al 31 dicembre: 2020-2024 dal report del Garante (tabella 10), 2025 dal DGMC
PRESENTI_31_12 = {'2020': 278, '2021': 318, '2022': 381, '2023': 495, '2024': 543, '2025': 572}


def comune_e_codice(df, col='IPM'):
    pos = df.columns.get_loc(col) + 1
    df.insert(pos, 'Comune', df[col].map(lambda s: COMUNE.get(s, '')))
    df.insert(pos + 1, 'Codice Istat comune', df['Comune'].map(lambda c: lookup_codice(c, 'comuni') if c else ''))
    df[col] = df[col].map(lambda s: NOME.get(s, s))
    return df


def it(d):
    return pd.Timestamp(d).strftime('%d/%m/%Y')


riep = pd.read_csv(f'{OUT_DIR}/riepilogo.csv', parse_dates=['data'])
sedi = pd.read_csv(f'{OUT_DIR}/sedi.csv', parse_dates=['data'], keep_default_na=False)
eta = pd.read_csv(f'{OUT_DIR}/eta_sede.csv', parse_dates=['data'], keep_default_na=False)
naz = pd.read_csv(f'{OUT_DIR}/eta_naz.csv', parse_dates=['data'])
pos = pd.read_csv(f'{OUT_DIR}/posizione.csv', parse_dates=['data'])
ultima = sedi['data'].max()
sheets = {}

# 1. Totale nazionale quindicinale
nt = naz[naz.eta == 'Totale'].set_index('data')
mino = naz[naz.eta.isin(['14-15', '16-17'])].groupby('data').tot.sum()
pt = pos[pos.gruppo == 'Totale'].pivot(index='data', columns='posizione', values='tot')
non_def = pt[['In attesa di 1° giudizio', 'Appellanti', 'Ricorrenti', 'Mista senza definitivo']].sum(axis=1)
r = riep.set_index('data')
naz_df = pd.DataFrame({
    'Data': r.index,
    'Detenuti presenti negli IPM': r.ipm, 'di cui maschi': r.ipm_m, 'di cui femmine': r.ipm_f,
    'Minorenni (14-17 anni)': mino, 'Giovani adulti (18-24 anni)': r.ipm - mino,
    'Italiani': nt.ita, 'Stranieri': nt.str, 'Stranieri (%)': (nt.str / nt.tot * 100).round(1),
    'In attesa di primo giudizio': pt['In attesa di 1° giudizio'],
    'Senza condanna definitiva': non_def,
    'Senza condanna definitiva (%)': (non_def / pt['Totale'] * 100).round(1),
    'Presenti nei Centri di prima accoglienza': r.cpa, 'Collocati in comunità ministeriali': r.com_min,
    'Collocati in comunità private': r.com_priv, 'Totale presenti nei servizi residenziali': r.tot_res,
    'In carico agli Uffici di servizio sociale per i minorenni': r.ussm,
}).reset_index(drop=True)
sheets['Totale nazionale'] = naz_df

# 2. Presenze per istituto (formato lungo)
lungo = sedi[sedi.ipm != 'Totale'].merge(eta[eta.ipm != 'Totale'], on=['data', 'ipm'], how='left')
lungo['ord'] = lungo.ipm.map(ORDINE.index)
lungo = lungo.sort_values(['data', 'ord'])
lungo_df = pd.DataFrame({
    'Data': lungo.data, 'IPM': lungo.ipm, 'Detenuti presenti': lungo.presenti,
    'Presenti 14-15 anni': lungo.eta_14_15, 'Presenti 16-17 anni': lungo.eta_16_17,
    'Presenti 18-20 anni': lungo.eta_18_20, 'Presenti 21-24 anni': lungo.eta_21_24,
    "Presenza media giornaliera dall'inizio dell'anno": lungo.presenza_media,
    "Ingressi dall'inizio dell'anno (esclusi trasferimenti)": lungo.ingressi_stabili,
    "Trasferimenti da altro IPM dall'inizio dell'anno": lungo.trasferimenti_da_ipm,
    "Ingressi totali dall'inizio dell'anno": lungo.ingressi_totali,
})
sheets['Presenze per istituto'] = comune_e_codice(lungo_df.reset_index(drop=True))

# 3. Presenti per istituto, una colonna per rilevazione
wide = sedi[sedi.ipm != 'Totale'].pivot(index='ipm', columns='data', values='presenti')
wide = wide.reindex([s for s in ORDINE if s in wide.index])
wide.columns = [it(c) for c in wide.columns]
wide = wide.reset_index().rename(columns={'ipm': 'IPM'})
sheets['Presenti serie'] = comune_e_codice(wide)

# 4. Affollamento all'ultima rilevazione
s26 = sedi[(sedi.ipm != 'Totale') & (sedi.data.dt.year == 2026)]
righe = []
for ipm in ORDINE:
    if ipm not in CAP:
        continue
    cap_alt, fonte_alt, nota_alt = CAP[ipm]
    if ipm in CAP_DGMC:
        cap, fonte, nota = CAP_DGMC[ipm], GARANTE, NOTE_DGMC.get(ipm, '')
    else:
        cap, fonte, nota = cap_alt, fonte_alt, nota_alt
    d = s26[s26.ipm == ipm]
    ult = d[d.data == ultima]
    pres = int(ult.presenti.iloc[0]) if len(ult) else None
    dic = sedi[(sedi.ipm == ipm) & (sedi.data == '2025-12-31')]
    mx = d.loc[d.presenti.idxmax()] if len(d) else None
    altra = ipm in CAP_DGMC and fonte_alt != GARANTE
    righe.append({
        'IPM': ipm, 'Capienza usata (posti)': cap, 'Fonte capienza usata': fonte, 'Note capienza usata': nota,
        'Presenti 31/12/2025': int(dic.presenti.iloc[0]) if len(dic) else None,
        f'Presenti {it(ultima)}': pres,
        f'Affollamento {it(ultima)} (%)': round(pres / cap * 100, 1) if pres is not None else None,
        'Massimo presenti 2026': int(mx.presenti) if mx is not None else None,
        'Data del massimo': mx.data if mx is not None else None,
        'Affollamento al massimo 2026 (%)': round(mx.presenti / cap * 100, 1) if mx is not None else None,
        f'Presenza media giornaliera 01/01-{it(ultima)[:5]}/2026': float(ult.presenza_media.iloc[0]) if len(ult) else None,
        'Capienza da altra fonte (posti)': cap_alt if altra else None,
        'Altra fonte meno DGMC (posti)': cap_alt - cap if altra else None,
        'Altra fonte': fonte_alt if altra else '', 'Note altra fonte': nota_alt if altra else '',
    })
aff = pd.DataFrame(righe)
sheets['Affollamento per istituto'] = comune_e_codice(aff)

# 5. Serie storica 2016-2025 (+2026 parziale)
si = pd.read_csv(f'{OUT_DIR}/serie_ingressi.csv')
sp = pd.read_csv(f'{OUT_DIR}/serie_presenza_media.csv')
serie = pd.DataFrame({
    'Anno': si.voce.astype(str), 'Ingressi negli IPM': si.tot, 'Ingressi di italiani': si.ita,
    'Ingressi di stranieri': si.str, 'Ingressi di femmine': si.tot_f,
    'Presenza media giornaliera': sp.tot, 'Presenza media giornaliera italiani': sp.ita,
    'Presenza media giornaliera stranieri': sp.str,
    'Presenti al 31/12': si.voce.astype(str).map(PRESENTI_31_12), 'Note': '',
})
serie.loc[serie.Anno.isin(['2020', '2021', '2022', '2023', '2024']), 'Note'] = \
    'Presenti al 31/12 dal report del Garante nazionale al 31/03/2025'
tot_ult = sedi[(sedi.ipm == 'Totale') & (sedi.data == ultima)].iloc[0]
serie.loc[len(serie)] = {'Anno': '2026', 'Ingressi negli IPM': int(tot_ult.ingressi_stabili),
                         'Presenza media giornaliera': float(tot_ult.presenza_media),
                         'Note': f'Dati parziali e provvisori, 01/01-{it(ultima)}'}
sheets['Serie storica 2016-2026'] = serie

# 6. Posizione giuridica
pl = pos.copy()
pl['ord'] = pl.posizione.map(['In attesa di 1° giudizio', 'Appellanti', 'Ricorrenti', 'Mista senza definitivo',
                              'Definitivi senza altri procedimenti a carico',
                              'Definitivi con altri procedimenti a carico', 'Totale'].index)
pl['g'] = pl.gruppo.map(['Minorenni', 'Giovani adulti', 'Totale'].index)
pl = pl.sort_values(['data', 'g', 'ord'])
sheets['Posizione giuridica'] = pd.DataFrame({'Data': pl.data, 'Gruppo': pl.gruppo, 'Posizione giuridica': pl.posizione,
                                              'Italiani': pl.ita, 'Stranieri': pl.str, 'Totale': pl.tot})


# 7-9. Confronto 2025 / 2026 per motivi di ingresso, delitti, uscite
def confronto(nome, label_sez):
    a = pd.read_csv(f'{OUT_2025}/{nome}.csv', keep_default_na=False)
    b = pd.read_csv(f'{OUT_DIR}/{nome}.csv', keep_default_na=False)
    for x in (a, b):
        x.loc[x.voce.str.startswith('Totale'), 'sezione'] = 'Totale'
        x['voce'] = x.voce.str.replace('di cui: ', '', regex=False)
    m = a[['sezione', 'voce', 'tot']].merge(b[['sezione', 'voce', 'ita', 'str', 'tot']], on=['sezione', 'voce'],
                                            how='outer', suffixes=('_2025', '_2026'), sort=False)
    # ordine: quello dell'anno 2025 (più voci), le voci solo 2026 subito dopo la riga che le precede
    ka, kb = list(zip(a.sezione, a.voce)), list(zip(b.sezione, b.voce))
    def ordine(r):
        k = (r.sezione, r.voce)
        if r.sezione == 'Totale':
            return len(ka) + 1
        if k in ka:
            return ka.index(k)
        prec = kb[kb.index(k) - 1]
        return (ka.index(prec) if prec in ka else len(ka)) + 0.5
    m['o'] = m.apply(ordine, axis=1)
    m = m.sort_values('o', kind='stable')
    nota = m.apply(lambda r: 'Voce non dettagliata dalla fonte in questo periodo (frequenza sotto la soglia di pubblicazione)'
                   if pd.isna(r.tot_2025) or pd.isna(r.tot_2026) else '', axis=1)
    return pd.DataFrame({label_sez: m.sezione, 'Voce': m.voce, 'Anno 2025': m.tot_2025,
                         f'2026 fino al {it(ultima)}': m.tot_2026,
                         f'2026 fino al {it(ultima)} italiani': m.ita, f'2026 fino al {it(ultima)} stranieri': m.str,
                         'Note': nota})


sheets['Ingressi per motivo'] = confronto('motivi_ingresso', 'Titolo di ingresso')
sheets['Delitti degli entrati'] = confronto('delitti_ingressi', 'Categoria')
sheets['Uscite per motivo'] = confronto('uscite', 'Tipo di uscita')

# 10. Provenienza all'ultima rilevazione
pv = pd.read_csv(f'{OUT_DIR}/provenienza.csv')
pv['voce'] = pv.voce.str.replace('secondo la provenienza e il sesso. ', '', regex=False) \
    .str.replace('di cui: ', '', regex=False)
aree = {'Italia', "Altri Paesi dell'Unione Europea", 'Altri Paesi europei', 'Africa', 'America', 'Asia', 'Oceania',
        'Totale'}
pv['Livello'] = pv.voce.map(lambda v: 'Area' if v in aree else 'Paese')
sheets[f'Provenienza {it(ultima)[:5].replace("/", "-")}'] = pd.DataFrame(
    {'Livello': pv.Livello, 'Provenienza': pv.voce, 'Maschi': pv.m, 'Femmine': pv.f, 'Totale': pv.tot})

# --- scrittura
HF, DF = Font(name='Arial', size=10, bold=True), Font(name='Arial', size=10)
L, R = Alignment(horizontal='left'), Alignment(horizontal='right')
wb = Workbook()
wb.remove(wb.active)
for name, df in sheets.items():
    ws = wb.create_sheet(name[:31])
    cols = list(df.columns)
    fmt = {}
    for c in cols:
        s = df[c]
        if c in ('Codice Istat comune', 'Anno'):
            fmt[c] = '@'
        elif pd.api.types.is_datetime64_any_dtype(s) or c.startswith('Data'):
            fmt[c] = 'DD/MM/YYYY'
        elif pd.api.types.is_numeric_dtype(s):
            v = s.dropna()
            fmt[c] = '0' if (v % 1 == 0).all() else ('0.0' if (v == v.round(1)).all() else '0.00')
        else:
            fmt[c] = None
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = HF
        cell.alignment = R if fmt[c] in ('0', '0.0', '0.00', 'DD/MM/YYYY') else L
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, (c, v) in enumerate(zip(cols, row), 1):
            if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT or v == '':
                v = None
            elif fmt[c] == 'DD/MM/YYYY':
                v = pd.Timestamp(v).to_pydatetime()
            elif fmt[c] == '0':
                v = int(v)
            elif fmt[c] in ('0.0', '0.00'):
                v = float(v)
            elif fmt[c] == '@':
                v = str(v)
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = DF
            if fmt[c]:
                cell.number_format = fmt[c]
            cell.alignment = R if fmt[c] in ('0', '0.0', '0.00', 'DD/MM/YYYY') else L
    ws.freeze_panes = 'A2'
    for j, c in enumerate(cols, 1):
        w = max([len(str(c))] + [len(str(x)) for x in df[c].astype(str).head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, w + 2), 60)
wb.save(DEST)
for k, v in sheets.items():
    print(k, v.shape)
print(aff[['IPM', 'Capienza usata (posti)', f'Presenti {it(ultima)}', f'Affollamento {it(ultima)} (%)',
           'Massimo presenti 2026', 'Data del massimo', 'Affollamento al massimo 2026 (%)']].to_string())
print(naz_df.to_string())
